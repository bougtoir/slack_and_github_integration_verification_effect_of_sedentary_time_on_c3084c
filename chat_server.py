import asyncio
import json
import os
import re
import shutil
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import requests
from docx import Document
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from paper_sandbox.ai_client import AIClient
from paper_sandbox.config import Config
from paper_sandbox.docx_writer import ManuscriptWriter
from paper_sandbox.figure_generator import FigureGenerator
from paper_sandbox.repo_publisher import publish
from paper_sandbox.slack import SlackNotifier
from paper_sandbox.stages import checks, journal_select

ROOT = Path(__file__).resolve().parent
STATIC = ROOT / "static"
SESSIONS_ROOT = ROOT / "workspace" / "sessions"
REPO = ROOT

BRAINSTORM_PROMPT = (
    "You are a research collaborator helping a clinician or researcher develop a paper idea. "
    "Use Socratic questioning and evidence-grounded suggestions to deepen the question. "
    "Ask concise follow-up questions (1-3) that clarify the population, exposure/outcome, "
    "data availability, and feasible analysis plan. Suggest relevant concepts or literature "
    "only if they are verifiable. "
    "When the user is ready, they can press the 'Start pipeline' button or type a command "
    "such as 'start pipeline'. "
    "At the end of every reply, include a clearly marked section 'Draft Background/Introduction' "
    "(2-3 paragraphs) that can be used directly as the manuscript Background/Introduction once "
    "the user is satisfied. Do not write the full paper now."
)

REVISION_PROMPT = (
    "You are an academic editor assisting with peer-review revisions. "
    "Given an original manuscript and reviewer comments, produce a revised manuscript JSON "
    "and a point-by-point response letter. "
    "Do not invent new empirical numbers; preserve all data-backed numbers from the original. "
    "If a requested change cannot be addressed without new data, state so honestly.\n\n"
    "Return ONLY a JSON object with this structure:\n"
    '{\n'
    '  "revised_manuscript": {\n'
    '    "title": "...",\n'
    '    "abstract": "...",\n'
    '    "sections": {"introduction":"...", "methods":"...", "results":"...", "discussion":"...", "conclusion":"..."},\n'
    '    "figures": [{"id":1, "caption":"..."}],\n'
    '    "tables": [{"id":1, "caption":"...", "headers":["..."], "rows":[["..."]]]\n'
    '  },\n'
    '  "response_letter": "..."\n'
    '}'
)

app = FastAPI()
app.mount("/static", StaticFiles(directory=str(STATIC)), name="static")

SESSIONS_ROOT.mkdir(parents=True, exist_ok=True)

# Serve per-session output files
@app.get("/output/{session_id}/{path:path}")
async def session_output(session_id: str, path: str):
    session = sessions.get(session_id)
    if session:
        output_dir = session.output_dir
    else:
        output_dir = SESSIONS_ROOT / session_id / "output"
    target = output_dir / path
    try:
        target.relative_to(output_dir)
    except ValueError:
        return JSONResponse({"ok": False, "error": "Invalid path"}, status_code=403)
    if not target.is_file():
        return JSONResponse({"ok": False, "error": "File not found"}, status_code=404)
    return FileResponse(target)


cfg = Config(str(ROOT / ".env"))


class ChatMessage(BaseModel):
    message: str
    session_id: Optional[str] = None
    start_pipeline: bool = False


class CandidateRequest(BaseModel):
    session_id: str


class ContinueRequest(BaseModel):
    session_id: str
    selected_index: int


class PeerReviewRequest(BaseModel):
    session_id: Optional[str] = None
    original_session_id: Optional[str] = None
    repo_url: Optional[str] = None
    reviewer_comments_text: Optional[str] = None
    decision: Optional[str] = "major"
    author_response: Optional[str] = None


@dataclass
class ChatSession:
    sid: str
    history: list = field(default_factory=list)
    previous_response_id: Optional[str] = None
    data_file: Optional[str] = None
    data_filename: Optional[str] = None
    protocol_files: list = field(default_factory=list)
    protocol_text: str = ""
    candidates: list = field(default_factory=list)
    output_dir: Path = field(default_factory=Path)
    upload_dir: Path = field(default_factory=Path)

    def __post_init__(self):
        self.upload_dir = SESSIONS_ROOT / self.sid / "uploads"
        self.output_dir = SESSIONS_ROOT / self.sid / "output"
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self.output_dir.mkdir(parents=True, exist_ok=True)


@dataclass
class PipelineJob:
    topic: str
    session: Optional[ChatSession]
    is_peer_review: bool = False
    task: Optional[asyncio.Task] = None
    done: bool = False
    error: Optional[str] = None
    summary: dict = field(default_factory=dict)


sessions: dict[str, ChatSession] = {}
jobs: dict[str, PipelineJob] = {}


def get_or_create_session(sid: Optional[str] = None) -> ChatSession:
    if sid and sid in sessions:
        return sessions[sid]
    new_sid = sid or os.urandom(16).hex()
    session = ChatSession(sid=new_sid)
    sessions[new_sid] = session
    return session


def _extract_text(response):
    text = getattr(response, "output_text", "") or ""
    if text:
        return text
    for item in getattr(response, "output", []) or []:
        if item.get("type") == "message" and item.get("role") == "assistant":
            for block in item.get("content", []):
                if block.get("type") in ("output_text", "text"):
                    return block.get("text", "")
    return ""


def _find_first_user_question(session: ChatSession) -> str:
    for entry in session.history:
        if entry.get("role") == "user":
            return entry.get("content", "")
    return ""


def _find_last_assistant_reply(session: ChatSession) -> Optional[str]:
    for entry in reversed(session.history):
        if entry.get("role") == "assistant":
            return entry.get("content", "")
    return None


def _extract_background(reply: str) -> str:
    m = re.search(r"(?:###\s*Draft Background/Introduction|Draft Background/Introduction)[\s:]*\n?(.*)", reply, re.S | re.I)
    if m:
        return m.group(1).strip()
    return reply.strip()


def _resolve_topic_and_background(session: ChatSession) -> tuple[str, str]:
    topic = _find_first_user_question(session)
    last = _find_last_assistant_reply(session) or ""
    background = session.protocol_text or _extract_background(last)
    return topic, background


def _container_upload_path(session_id: str, filename: str) -> str:
    return f"/app/workspace/sessions/{session_id}/uploads/{filename}"


def _extract_docx_text(path: Path) -> Optional[str]:
    try:
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text).strip()
    except Exception:
        return None


def _extract_text_file(path: Path) -> Optional[str]:
    try:
        return path.read_text(encoding="utf-8", errors="ignore").strip()
    except Exception:
        return None


def _classify_file(path: Path) -> str:
    suffix = path.suffix.lower()
    data_suffixes = {".csv", ".tsv", ".xlsx", ".xls", ".parquet", ".json", ".sas7bdat", ".dta", ".rds"}
    text_suffixes = {".txt", ".md", ".docx"}
    if suffix in data_suffixes:
        if suffix == ".json":
            try:
                data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
                if isinstance(data, list) and data:
                    return "data"
            except Exception:
                pass
        return "data"
    if suffix in text_suffixes:
        return "protocol"
    return "unknown"


def _handle_uploaded_file(session: ChatSession, upload: UploadFile) -> dict:
    safe_name = Path(upload.filename).name
    dest = session.upload_dir / safe_name
    content = upload.file.read() if hasattr(upload.file, "read") else upload.read()
    if isinstance(content, str):
        content = content.encode("utf-8")
    dest.write_bytes(content)
    kind = _classify_file(dest)
    result = {"filename": safe_name, "path": str(dest), "container_path": _container_upload_path(session.sid, safe_name), "kind": kind}
    if kind == "data":
        session.data_file = _container_upload_path(session.sid, safe_name)
        session.data_filename = safe_name
    elif kind == "protocol":
        session.protocol_files.append(safe_name)
        text = _extract_docx_text(dest) if dest.suffix.lower() == ".docx" else _extract_text_file(dest)
        if text:
            if session.protocol_text:
                session.protocol_text += "\n\n" + text
            else:
                session.protocol_text = text
        result["extracted_text"] = bool(text)
    return result


async def brainstorm(message: str, session: ChatSession) -> tuple[str, str]:
    try:
        from perplexity import Perplexity
        client = Perplexity(api_key=cfg.perplexity_api_key)
        params = {
            "preset": cfg.perplexity_preset,
            "input": message,
            "instructions": BRAINSTORM_PROMPT,
            "tools": [{"type": "web_search", "search_context_size": "medium"}],
            "store": True,
        }
        if session.previous_response_id:
            params["previous_response_id"] = session.previous_response_id
        response = await asyncio.to_thread(client.responses.create, **params)
        text = _extract_text(response)
        if response and getattr(response, "id", None):
            session.previous_response_id = response.id
        return text or "(No response from Perplexity)", response.id if response else ""
    except Exception as e:
        return f"Brainstorm failed: {e}", ""


def _build_pipeline_input(topic: str, background: Optional[str], session: Optional[ChatSession], chosen_journal: Optional[dict] = None) -> dict:
    input_data: dict = {"topic": topic, "authors": ["Sandbox Author"]}
    if session and session.data_file:
        input_data["data_file"] = session.data_file
    protocol = session.protocol_text if session and session.protocol_text else background
    if protocol:
        input_data["protocol"] = protocol
    if chosen_journal:
        input_data["chosen_journal"] = chosen_journal
    return input_data


async def _run_cli(session: ChatSession, input_data: dict, mode: str = "run") -> dict:
    payload = json.dumps(input_data, ensure_ascii=False)
    session.output_dir.mkdir(parents=True, exist_ok=True)
    local_input = session.output_dir / "input.json"
    local_input.write_text(payload, encoding="utf-8")
    summary_path = session.output_dir / "summary.json"

    env_output_dir = f"/app/workspace/sessions/{session.sid}"
    env_args = [
        "-e", f"OUTPUT_DIR={env_output_dir}",
        "-e", f"WORKSPACE=/app/workspace",
    ]
    for key in ("SLACK_BOT_TOKEN", "SLACK_WEBHOOK_URL", "GITHUB_REPO_CREATE_TOKEN"):
        val = os.getenv(key)
        if val:
            env_args.extend(["-e", f"{key}={val}"])

    use_docker = False
    if shutil.which("docker"):
        check = await asyncio.create_subprocess_exec(
            "docker", "compose", "ps", "--services", "--filter", "status=running",
            cwd=str(REPO),
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        stdout, _ = await check.communicate()
        use_docker = "sandbox" in stdout.decode("utf-8", errors="ignore")

    if use_docker:
        container_input = f"{env_output_dir}/input.json"
        cmd = [
            "docker", "compose", "exec", "-T", *env_args,
            "sandbox", "sh", "-c",
            f"mkdir -p {env_output_dir} && cat > {container_input} && python -m paper_sandbox.cli --mode {mode} --input {container_input} --output-summary {env_output_dir}/summary.json",
        ]
        proc = await asyncio.create_subprocess_exec(
            *cmd,
            cwd=str(REPO),
            stdin=asyncio.subprocess.PIPE,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        stdout, stderr = await proc.communicate(payload.encode("utf-8"))
    else:
        env = os.environ.copy()
        env["OUTPUT_DIR"] = str(session.output_dir)
        env["WORKSPACE"] = str(SESSIONS_ROOT)
        env["GIT_AUTO_COMMIT"] = "false"
        cmd = [
            sys.executable, "-m", "paper_sandbox.cli",
            "--mode", mode,
            "--input", str(local_input),
            "--output-summary", str(summary_path),
        ]
        proc = await asyncio.create_subprocess_exec(
            *cmd,
            cwd=str(REPO),
            env=env,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        stdout, stderr = await proc.communicate()

    if proc.returncode != 0:
        raise RuntimeError(stderr.decode("utf-8", errors="ignore")[:2000])
    if summary_path.exists():
        return json.loads(summary_path.read_text(encoding="utf-8"))
    return {}


async def run_pipeline(topic: str, background: Optional[str] = None, session: Optional[ChatSession] = None, chosen_journal: Optional[dict] = None):
    mode = "continue" if chosen_journal else "run"
    input_data = _build_pipeline_input(topic, background, session, chosen_journal)
    return await _run_cli(session or get_or_create_session(), input_data, mode=mode)


async def _publish_open_repo(topic: str, session: ChatSession) -> tuple[Optional[str], Optional[str], Optional[str]]:
    return await asyncio.to_thread(publish, topic, str(REPO), str(session.output_dir))


def _notify_repo_created(topic: str, repo_name: str, repo_url: str, chosen_journal: Optional[dict]):
    try:
        notifier = SlackNotifier()
        journal_name = chosen_journal.get("name") if chosen_journal else None
        message = f"Paper Sandbox: submission prepared\nTopic: {topic}"
        if journal_name:
            message += f"\nJournal: {journal_name}"
        message += f"\nRepo: {repo_url}"
        notifier.create_and_post(repo_name, message)
    except Exception:
        pass


def _build_output_url(session: ChatSession, rel_path: str) -> str:
    return f"/output/{session.sid}/{rel_path}"


def _list_output_links(session: ChatSession, summary: dict) -> dict:
    out = session.output_dir
    links: dict = {}
    manifest_path = out / "manifest.json"
    for key, name in [
        ("manuscript_url", "manuscript.docx"),
        ("cover_letter_url", "cover_letter.docx"),
        ("review_report_url", "reviewer_review.docx"),
        ("validation_report_url", "validation_report.docx"),
        ("tables_url", "tables.docx"),
        ("revised_manuscript_url", "revised_manuscript.docx"),
        ("response_letter_url", "response_to_reviewers.docx"),
        ("journal_candidates_url", "journal_candidates.md"),
        ("journal_table_url", "journal_table.md"),
        ("pre_submission_checklist_url", "pre_submission_checklist.docx"),
        ("pre_submission_checklist_md_url", "pre_submission_checklist.md"),
    ]:
        if (out / name).exists():
            links[key] = _build_output_url(session, name)

    if (out / "summary.json").exists():
        links["summary_url"] = _build_output_url(session, "summary.json")
    if (out / "package.zip").exists():
        links["package_url"] = _build_output_url(session, "package.zip")

    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            figures = manifest.get("figures") or []
            figure_urls = []
            for p in figures:
                pp = Path(p)
                if pp.parent.name == "figures":
                    rel = f"figures/{pp.name}"
                else:
                    rel = pp.name
                if (out / rel).exists():
                    figure_urls.append(_build_output_url(session, rel))
            links["figure_urls"] = figure_urls
        except Exception:
            pass

    links["repo_url"] = summary.get("repo_url")
    links["repo_error"] = summary.get("repo_error")
    return links


async def pipeline_worker(job_id: str, topic: str, background: Optional[str], session: ChatSession, chosen_journal: Optional[dict] = None):
    job = jobs.get(job_id)
    if not job:
        return
    try:
        summary = await run_pipeline(topic, background, session, chosen_journal)
        repo_url, repo_name, repo_error = await _publish_open_repo(topic, session)
        if repo_url:
            summary["repo_url"] = repo_url
        if repo_name:
            summary["repo_name"] = repo_name
        if repo_error:
            summary["repo_error"] = repo_error
        job.summary = summary
        summary_path = session.output_dir / "summary.json"
        summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
        if repo_url and repo_name:
            _notify_repo_created(topic, repo_name, repo_url, chosen_journal)
    except Exception as e:
        job.error = str(e)
    finally:
        job.done = True


def _parse_manuscript_json(text: str) -> dict:
    try:
        m = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.DOTALL)
        if m:
            return json.loads(m.group(1))
        m = re.search(r"(\{.*\})", text, re.DOTALL)
        if m:
            return json.loads(m.group(1))
    except Exception:
        pass
    return {}


def _docx_to_text(path: Path) -> str:
    text = _extract_docx_text(path)
    return text or ""


def _revision_worker(job_id: str, session: ChatSession, original_docx_path: Path, comments_text: str, decision: str, author_response: Optional[str], repo_url: Optional[str]):
    job = jobs.get(job_id)
    if not job:
        return
    try:
        client = AIClient(cfg.deepseek_api_key, cfg.deepseek_base_url, cfg.deepseek_model)
        original_text = _docx_to_text(original_docx_path)
        prompt = (
            f"{REVISION_PROMPT}\n\n"
            f"Decision: {decision}\n"
            f"Original manuscript:\n{original_text[:12000]}\n\n"
            f"Reviewer comments:\n{comments_text[:12000]}\n\n"
        )
        if author_response:
            prompt += f"Authors' intended response / additional context:\n{author_response}\n\n"
        if repo_url:
            prompt += f"Original submission repository: {repo_url}\n"

        response = client.chat(prompt, temperature=0.4)
        parsed = _parse_manuscript_json(response)

        revised = parsed.get("revised_manuscript") or {"title": "Revised manuscript", "abstract": "", "sections": {"introduction":"", "methods":"", "results":"", "discussion":"", "conclusion":""}, "figures": [], "tables": []}
        response_letter = parsed.get("response_letter") or response[:5000] if response else "No response generated."

        writer = ManuscriptWriter(cfg)
        fg = FigureGenerator(session.output_dir / "figures")
        figure_paths = []
        if revised.get("figures"):
            for fig in revised["figures"]:
                name = f"figure_{fig.get('id', 1)}"
                png, tiff, pptx = fg.demo_figure(fig.get("caption", "Figure"), name=name)
                figure_paths.extend([png, tiff, pptx])

        tables_path, table_pptx = fg.table_docx(revised.get("tables", []))
        if table_pptx:
            figure_paths.append(table_pptx)

        revised_docx = writer.write_manuscript(revised, session.output_dir, figure_paths=figure_paths, table_path=tables_path, filename="revised_manuscript.docx")

        response_doc = Document()
        response_doc.add_heading("Response to Reviewers", level=1)
        for para in str(response_letter).split("\n"):
            response_doc.add_paragraph(para)
        response_letter_path = session.output_dir / "response_to_reviewers.docx"
        response_doc.save(response_letter_path)

        check_results = checks.run_all_checks(revised, language=cfg.language)
        (session.output_dir / "revision_checks.json").write_text(json.dumps(check_results, indent=2, ensure_ascii=False), encoding="utf-8")

        summary = {
            "stage": "peer_review",
            "revised_manuscript": _build_output_url(session, revised_docx.name),
            "response_letter": _build_output_url(session, response_letter_path.name),
            "original_repo_url": repo_url,
        }
        (session.output_dir / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
        job.summary = summary
    except Exception as e:
        job.error = str(e)
    finally:
        job.done = True


@app.get("/")
async def root():
    return FileResponse(str(STATIC / "index.html"))


@app.post("/upload")
async def upload(
    session_id: Optional[str] = Form(None),
    files: list[UploadFile] = File(default=[]),
    data: Optional[UploadFile] = File(None),
    protocol: Optional[UploadFile] = File(None),
):
    session = get_or_create_session(session_id)
    results = []

    # Backward-compatible single-file fields
    legacy = []
    if data and data.filename:
        legacy.append(data)
    if protocol and protocol.filename:
        legacy.append(protocol)
    if files:
        legacy.extend(files)

    for upload in legacy:
        results.append(_handle_uploaded_file(session, upload))

    return JSONResponse({"ok": True, "session_id": session.sid, "files": results})


@app.post("/chat")
async def chat(msg: ChatMessage):
    session = get_or_create_session(msg.session_id)
    session.history.append({"role": "user", "content": msg.message})

    if msg.start_pipeline:
        topic, background = _resolve_topic_and_background(session)
        if not topic:
            topic = msg.message
        job_id = os.urandom(12).hex()
        job = PipelineJob(topic=topic, session=session)
        jobs[job_id] = job
        job.task = asyncio.create_task(pipeline_worker(job_id, topic, background, session))
        return JSONResponse({
            "ok": True,
            "stage": "started",
            "session_id": session.sid,
            "job_id": job_id,
            "topic": topic,
            "message": "Pipeline started. It will continue in the background; poll /status/{job_id}.",
        })

    reply, _ = await brainstorm(msg.message, session)
    session.history.append({"role": "assistant", "content": reply})
    return JSONResponse({
        "ok": True,
        "stage": "brainstorm",
        "session_id": session.sid,
        "reply": reply,
    })


@app.post("/journal_candidates")
async def journal_candidates(req: CandidateRequest):
    session = sessions.get(req.session_id)
    if not session:
        return JSONResponse({"ok": False, "error": "Session not found"}, status_code=404)
    topic, _ = _resolve_topic_and_background(session)
    if not topic:
        return JSONResponse({"ok": False, "error": "No topic found in session"}, status_code=400)

    ranked, table_md = await asyncio.to_thread(journal_select.select_journals, cfg, topic, False)
    session.output_dir.mkdir(parents=True, exist_ok=True)
    (session.output_dir / "journal_candidates.md").write_text(table_md, encoding="utf-8")
    session.candidates = ranked[:10]

    candidates = [
        {"index": i, "name": j["name"], "if_2023": j["if_2023"], "apc_usd": j["apc_usd"], "hybrid": j["hybrid"], "publisher": j["publisher"]}
        for i, j in enumerate(ranked[:10])
    ]
    return JSONResponse({
        "ok": True,
        "stage": "candidates",
        "session_id": session.sid,
        "topic": topic,
        "candidates": candidates,
        "markdown": table_md,
        "markdown_url": _build_output_url(session, "journal_candidates.md"),
    })


@app.post("/continue_pipeline")
async def continue_pipeline(req: ContinueRequest):
    session = sessions.get(req.session_id)
    if not session:
        return JSONResponse({"ok": False, "error": "Session not found"}, status_code=404)
    if req.selected_index < 0 or req.selected_index >= len(session.candidates):
        return JSONResponse({"ok": False, "error": f"Invalid index. Use 0-{len(session.candidates)-1}"}, status_code=400)

    chosen = session.candidates[req.selected_index]
    topic, background = _resolve_topic_and_background(session)
    if not topic:
        topic = chosen.get("name", "Untitled")

    job_id = os.urandom(12).hex()
    job = PipelineJob(topic=topic, session=session)
    jobs[job_id] = job
    job.task = asyncio.create_task(pipeline_worker(job_id, topic, background, session, chosen))
    return JSONResponse({
        "ok": True,
        "stage": "started",
        "session_id": session.sid,
        "job_id": job_id,
        "topic": topic,
        "chosen_journal": chosen["name"],
        "message": "Submission preparation started. Poll /status/{job_id}.",
    })


@app.post("/peer_review")
async def peer_review(
    session_id: Optional[str] = Form(None),
    original_session_id: Optional[str] = Form(None),
    repo_url: Optional[str] = Form(None),
    reviewer_comments_text: Optional[str] = Form(None),
    reviewer_comments: Optional[UploadFile] = File(None),
    manuscript: Optional[UploadFile] = File(None),
    decision: str = Form("major"),
    author_response: Optional[str] = Form(None),
):
    session = get_or_create_session(session_id)
    # Resolve original manuscript
    original_docx: Optional[Path] = None
    if manuscript and manuscript.filename:
        safe_name = Path(manuscript.filename).name
        dest = session.upload_dir / safe_name
        dest.write_bytes(await manuscript.read())
        original_docx = dest
    elif repo_url:
        # Try to fetch deliverables/manuscript.docx from a previously published public repo
        try:
            raw_url = repo_url.replace("github.com", "raw.githubusercontent.com").rstrip("/") + "/main/deliverables/manuscript.docx"
            r = requests.get(raw_url, timeout=60)
            r.raise_for_status()
            original_docx = session.upload_dir / "manuscript.docx"
            original_docx.write_bytes(r.content)
        except Exception:
            pass
    elif original_session_id:
        if original_session_id in sessions:
            original_session = sessions[original_session_id]
            candidate = original_session.output_dir / "manuscript.docx"
        else:
            candidate = SESSIONS_ROOT / original_session_id / "output" / "manuscript.docx"
        if candidate.exists():
            original_docx = candidate

    if not original_docx or not original_docx.exists():
        return JSONResponse({"ok": False, "error": "Original manuscript not found. Upload a .docx file, provide repo_url, or original_session_id."}, status_code=400)

    # Resolve reviewer comments
    comments_text = reviewer_comments_text or ""
    if reviewer_comments and reviewer_comments.filename:
        comments_path = session.upload_dir / Path(reviewer_comments.filename).name
        comments_path.write_bytes(await reviewer_comments.read())
        txt = _extract_docx_text(comments_path) if comments_path.suffix.lower() == ".docx" else _extract_text_file(comments_path)
        if txt:
            comments_text = txt
    if not comments_text:
        return JSONResponse({"ok": False, "error": "Reviewer comments required as text or file."}, status_code=400)

    job_id = os.urandom(12).hex()
    job = PipelineJob(topic="Peer review revision", session=session, is_peer_review=True)
    jobs[job_id] = job
    job.task = asyncio.create_task(
        asyncio.to_thread(
            _revision_worker, job_id, session, original_docx, comments_text, decision, author_response, repo_url
        )
    )
    return JSONResponse({
        "ok": True,
        "stage": "started",
        "session_id": session.sid,
        "job_id": job_id,
        "message": "Peer review revision started. Poll /status/{job_id}.",
    })


@app.get("/status/{job_id}")
async def status(job_id: str):
    job = jobs.get(job_id)
    if not job:
        return JSONResponse({"ok": False, "error": "Job not found"}, status_code=404)
    if not job.done:
        return JSONResponse({"ok": True, "stage": "running", "topic": job.topic, "session_id": job.session.sid if job.session else None})
    if job.error:
        return JSONResponse({"ok": False, "stage": "pipeline", "topic": job.topic, "error": job.error}, status_code=500)

    session = job.session
    links = _list_output_links(session, job.summary) if session else {}

    if job.is_peer_review:
        return JSONResponse({
            "ok": True,
            "stage": "peer_review",
            "session_id": session.sid if session else None,
            "topic": job.topic,
            **links,
            "message": "Peer review revision finished.",
        })

    chosen = job.summary.get("chosen_journal", {}) if isinstance(job.summary, dict) else {}
    if isinstance(chosen, str):
        chosen = {"name": chosen}
    return JSONResponse({
        "ok": True,
        "stage": "pipeline",
        "session_id": session.sid if session else None,
        "topic": job.topic,
        "journal": chosen.get("name", "Unknown"),
        "message": f"Pipeline finished. Chosen journal: {chosen.get('name', 'Unknown')}.",
        **links,
    })


@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    return JSONResponse({"ok": False, "error": str(exc)}, status_code=500)
